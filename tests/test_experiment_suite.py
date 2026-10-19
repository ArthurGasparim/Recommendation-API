import json

import pandas as pd
from openpyxl import load_workbook
from sqlmodel import Session

import experiments.real_runner as real_runner
from db.db import AdminJob, engine
from experiments.real_runner import RealExperimentSuiteRequest
from experiments.suite_report import build_suite_workbook
from tests.conftest import authenticate
from tests.test_real_experiment import measured_frames


def _metrics(value: float) -> dict:
    return {
        "ndcg_at_k": value,
        "precision_at_k": value - 0.10,
        "recall_at_k": value - 0.08,
        "hit_rate_at_k": value + 0.05,
        "catalog_coverage": value - 0.04,
        "latency_p50_ms": 10.0,
        "latency_p95_ms": 20.0,
        "latency_p99_ms": 25.0,
        "availability_percent": 99.5,
        "failed_requests": 1,
    }


def _quality(score: float) -> dict:
    return {
        "quality_score": score,
        "invalid_rate": 0.1 if score < 1 else 0.0,
        "orphan_rate": 0.05 if score < 1 else 0.0,
        "duplicate_rate": 0.05 if score < 1 else 0.0,
        "mean_rating": 4.0,
        "records": 100,
    }


def _fake_run(request, books_df, ratings_df, to_read_df, **kwargs) -> dict:
    measured_value = 0.74 if request.with_observability else 0.61
    return {
        "baseline": _metrics(0.80),
        "measured": _metrics(measured_value),
        "data_quality": {
            "baseline": _quality(1.0),
            "effective": _quality(
                0.97 if request.with_observability else 0.72
            ),
        },
        "observability": {
            "detected": request.with_observability,
            "response_applied": request.with_observability,
            "signal": "anomalia detectada" if request.with_observability else None,
            "action": "snapshot restaurado" if request.with_observability else "nenhuma ação",
            "detection_ms": 12.0 if request.with_observability else None,
            "recovery_ms": 18.0 if request.with_observability else None,
        },
        "attack": {"type": request.attack_type, "injected_records": 3},
        "timing": {"total_seconds": 0.02},
    }


def _suite_result(monkeypatch=None) -> dict:
    if monkeypatch is not None:
        monkeypatch.setattr(real_runner, "run_real_experiment", _fake_run)
    books, ratings, to_read = measured_frames()
    progress = []
    result = real_runner.run_real_experiment_suite(
        RealExperimentSuiteRequest(sample_users=5, top_k=5),
        books,
        ratings,
        to_read,
        lambda completed, total, message: progress.append(
            (completed, total, message)
        ),
    )
    assert progress[-1][0:2] == (36, 36)
    return result


def test_suite_runs_every_model_attack_and_observability_pair(monkeypatch):
    calls = []

    def tracked_run(request, *frames, **kwargs):
        calls.append((request.model, request.attack_type, request.with_observability))
        return _fake_run(request, *frames, **kwargs)

    monkeypatch.setattr(real_runner, "run_real_experiment", tracked_run)
    result = _suite_result()

    assert result["comparisons_count"] == 18
    assert result["scenario_runs"] == 36
    assert len(calls) == 36
    assert len(set(calls)) == 36
    assert result["method"]["database_mutated"] is False
    first = result["comparisons"][0]
    assert first["baseline"]["ndcg_at_k"] == 0.80
    assert first["without_observability"]["ndcg_at_k"] == 0.61
    assert first["with_observability"]["ndcg_at_k"] == 0.74


def test_xlsx_contains_quality_operation_and_methodology(monkeypatch):
    result = _suite_result(monkeypatch)
    workbook = load_workbook(build_suite_workbook(result), data_only=True)

    assert workbook.sheetnames == ["Qualidade", "Operação", "Dados", "Metodologia"]
    quality = workbook["Qualidade"]
    assert quality.max_row == 22
    assert quality["A4"].value == "Modelo"
    assert quality["C4"].value == "NDCG - Inicial"
    assert quality["D4"].value == "NDCG - Sem observabilidade"
    assert quality["E4"].value == "NDCG - Com observabilidade"
    assert quality["C5"].value == 0.80
    assert quality["D5"].value == 0.61
    assert quality["E5"].value == 0.74
    data_quality = workbook["Dados"]
    assert data_quality["A4"].value == "Modelo"
    assert data_quality["C4"].value == "Score de qualidade - Inicial"
    assert data_quality["D5"].value == 0.72
    assert data_quality["E5"].value == 0.97


def test_completed_suite_can_be_downloaded_again(client, monkeypatch):
    admin_headers = authenticate(client, 1)
    result = _suite_result(monkeypatch)
    result["suite_id"] = "suite-completed"
    with Session(engine) as session:
        session.add(
            AdminJob(
                id="suite-completed",
                kind="real_experiment_suite",
                status="completed",
                progress=100,
                message="Bateria concluída",
                requested_by=1,
                request_json="{}",
                result_json=json.dumps(result),
            )
        )
        session.add(
            AdminJob(
                id="suite-running",
                kind="real_experiment_suite",
                status="running",
                progress=50,
                message="Em andamento",
                requested_by=1,
                request_json="{}",
            )
        )
        session.commit()

    endpoint = "/admin/experiments/real/suite/suite-completed/xlsx"
    assert client.get(endpoint).status_code == 401
    response = client.get(endpoint, headers=admin_headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content.startswith(b"PK")
    assert "suite-completed.xlsx" in response.headers["content-disposition"]

    pending = client.get(
        "/admin/experiments/real/suite/suite-running/xlsx",
        headers=admin_headers,
    )
    assert pending.status_code == 409

    metrics = client.get("/admin/experiments/real/metrics", headers=admin_headers)
    assert metrics.status_code == 200
    latest = metrics.json()["latest_suite"]
    assert latest["id"] == "suite-completed"
    assert latest["summary"]["scenario_runs"] == 36
    assert latest["download_url"] == endpoint
