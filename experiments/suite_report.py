"""Geração do relatório XLSX da bateria de experimentos medidos."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "0B2440"
BLUE = "DDEBF7"
RED = "FCE4D6"
GREEN = "E2F0D9"
PURPLE = "E4DFEC"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"

QUALITY_METRICS = [
    ("ndcg_at_k", "NDCG"),
    ("precision_at_k", "Precision"),
    ("recall_at_k", "Recall"),
    ("hit_rate_at_k", "Hit rate"),
    ("catalog_coverage", "Cobertura"),
    ("data_quality_score", "Qualidade dos dados"),
]


def _quality_values(comparison: dict, metric: str) -> tuple[float, float, float]:
    if metric == "data_quality_score":
        source = comparison["data_quality"]
        return (
            float(source["baseline"]),
            float(source["without_observability"]),
            float(source["with_observability"]),
        )
    return (
        float(comparison["baseline"][metric]),
        float(comparison["without_observability"][metric]),
        float(comparison["with_observability"][metric]),
    )


def _prepare_sheet(sheet, title: str, subtitle: str, headers: list[str]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(color=WHITE, bold=True, size=16)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    subtitle_cell = sheet.cell(2, 1, subtitle)
    subtitle_cell.fill = PatternFill("solid", fgColor=LIGHT)
    subtitle_cell.font = Font(color="526477", italic=True, size=10)
    subtitle_cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 30

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(4, column, header)
        fill = BLUE
        if "Sem observabilidade" in header:
            fill = RED
        elif "Com observabilidade" in header:
            fill = GREEN
        elif "Δ" in header or "Diferença" in header:
            fill = PURPLE
        elif column <= 2:
            fill = NAVY
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color=WHITE if column <= 2 else "172033", bold=True, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[4].height = 42
    sheet.freeze_panes = "C5"


def _finish_table(sheet, table_name: str, max_column: int, max_row: int) -> None:
    if max_row >= 5:
        reference = f"A4:{get_column_letter(max_column)}{max_row}"
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    sheet.auto_filter.ref = f"A4:{get_column_letter(max_column)}{max_row}"
    sheet.column_dimensions["A"].width = 27
    sheet.column_dimensions["B"].width = 24
    for column in range(3, max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18


def _quality_sheet(workbook: Workbook, suite: dict) -> None:
    sheet = workbook.active
    sheet.title = "Qualidade"
    headers = ["Modelo", "Ataque"]
    for _, label in QUALITY_METRICS:
        headers.extend(
            [
                f"{label} - Inicial",
                f"{label} - Sem observabilidade",
                f"{label} - Com observabilidade",
                f"{label} - Δ com - sem",
            ]
        )
    _prepare_sheet(
        sheet,
        "Qualidade inicial × sem observabilidade × com observabilidade",
        (
            f"{suite['comparisons_count']} combinações e {suite['scenario_runs']} cenários "
            f"medidos em {suite['total_seconds']:.1f} segundos. Valores maiores são melhores."
        ),
        headers,
    )

    delta_columns: list[int] = []
    for row_index, comparison in enumerate(suite["comparisons"], 5):
        row: list[object] = [comparison["model_label"], comparison["attack_label"]]
        for metric, _ in QUALITY_METRICS:
            baseline, without, with_observability = _quality_values(comparison, metric)
            row.extend(
                [baseline, without, with_observability, with_observability - without]
            )
        sheet.append(row)
        for column in range(3, len(headers) + 1):
            sheet.cell(row_index, column).number_format = "0.000000"

    for metric_index in range(len(QUALITY_METRICS)):
        delta_column = 6 + metric_index * 4
        delta_columns.append(delta_column)
        if suite["comparisons"]:
            letter = get_column_letter(delta_column)
            sheet.conditional_formatting.add(
                f"{letter}5:{letter}{4 + len(suite['comparisons'])}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )
    _finish_table(sheet, "ComparacaoQualidade", len(headers), 4 + len(suite["comparisons"]))


def _operations_sheet(workbook: Workbook, suite: dict) -> None:
    sheet = workbook.create_sheet("Operação")
    headers = [
        "Modelo",
        "Ataque",
        "Latência p50 - Inicial (ms)",
        "Latência p50 - Sem observabilidade (ms)",
        "Latência p50 - Com observabilidade (ms)",
        "Latência p95 - Inicial (ms)",
        "Latência p95 - Sem observabilidade (ms)",
        "Latência p95 - Com observabilidade (ms)",
        "Latência p99 - Inicial (ms)",
        "Latência p99 - Sem observabilidade (ms)",
        "Latência p99 - Com observabilidade (ms)",
        "Disponibilidade - Inicial (%)",
        "Disponibilidade - Sem observabilidade (%)",
        "Disponibilidade - Com observabilidade (%)",
        "Falhas - Inicial",
        "Falhas - Sem observabilidade",
        "Falhas - Com observabilidade",
        "Ataque detectado",
        "Resposta aplicada",
        "Sinal",
        "Ação",
        "Detecção (ms)",
        "Recuperação (ms)",
        "Tempo sem observabilidade (s)",
        "Tempo com observabilidade (s)",
    ]
    _prepare_sheet(
        sheet,
        "Desempenho, disponibilidade e resposta observável",
        "Medições operacionais da mesma bateria; latências e tempos são valores observados durante cada execução.",
        headers,
    )
    for row_index, comparison in enumerate(suite["comparisons"], 5):
        baseline = comparison["baseline"]
        without = comparison["without_observability"]
        with_observability = comparison["with_observability"]
        observation = comparison["observability"]
        timing = comparison["timing"]
        sheet.append(
            [
                comparison["model_label"],
                comparison["attack_label"],
                baseline["latency_p50_ms"],
                without["latency_p50_ms"],
                with_observability["latency_p50_ms"],
                baseline["latency_p95_ms"],
                without["latency_p95_ms"],
                with_observability["latency_p95_ms"],
                baseline["latency_p99_ms"],
                without["latency_p99_ms"],
                with_observability["latency_p99_ms"],
                baseline["availability_percent"],
                without["availability_percent"],
                with_observability["availability_percent"],
                baseline["failed_requests"],
                without["failed_requests"],
                with_observability["failed_requests"],
                "Sim" if observation["detected"] else "Não",
                "Sim" if observation["response_applied"] else "Não",
                observation.get("signal") or "",
                observation.get("action") or "",
                observation.get("detection_ms"),
                observation.get("recovery_ms"),
                timing["without_observability_seconds"],
                timing["with_observability_seconds"],
            ]
        )
        for column in range(3, len(headers) + 1):
            if isinstance(sheet.cell(row_index, column).value, (int, float)):
                sheet.cell(row_index, column).number_format = "0.000"
    _finish_table(sheet, "ComparacaoOperacional", len(headers), 4 + len(suite["comparisons"]))
    sheet.column_dimensions["T"].width = 38
    sheet.column_dimensions["U"].width = 38


def _data_quality_sheet(workbook: Workbook, suite: dict) -> None:
    sheet = workbook.create_sheet("Dados")
    metrics = [
        ("quality_score", "Score de qualidade"),
        ("invalid_rate", "Taxa de domínio inválido"),
        ("orphan_rate", "Taxa de referências órfãs"),
        ("duplicate_rate", "Taxa de duplicidade"),
        ("mean_rating", "Nota média"),
        ("records", "Registros"),
    ]
    headers = ["Modelo", "Ataque"]
    for _, label in metrics:
        headers.extend(
            [
                f"{label} - Inicial",
                f"{label} - Sem observabilidade",
                f"{label} - Com observabilidade",
            ]
        )
    _prepare_sheet(
        sheet,
        "Degradação dos dados antes e depois da resposta",
        (
            "O cenário sem observabilidade mantém o lote atacado. O cenário com "
            "observabilidade mostra o estado efetivo após rollback ou quarentena."
        ),
        headers,
    )
    for row_index, comparison in enumerate(suite["comparisons"], 5):
        details = comparison["data_quality"]["details"]
        row: list[object] = [comparison["model_label"], comparison["attack_label"]]
        for metric, _ in metrics:
            row.extend(
                [
                    details["baseline"][metric],
                    details["without_observability"][metric],
                    details["with_observability"][metric],
                ]
            )
        sheet.append(row)
        for column in range(3, len(headers) + 1):
            value = sheet.cell(row_index, column).value
            if isinstance(value, float):
                sheet.cell(row_index, column).number_format = "0.000000"
    _finish_table(
        sheet,
        "ComparacaoQualidadeDados",
        len(headers),
        4 + len(suite["comparisons"]),
    )


def _method_sheet(workbook: Workbook, suite: dict) -> None:
    sheet = workbook.create_sheet("Metodologia")
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 95
    sheet.append(["Campo", "Valor"])
    configuration = suite["configuration"]
    rows = [
        ("Tipo", "Bateria de experimentos medidos"),
        ("Data da medição", suite["measured_at"]),
        ("Modelos", suite["models"]),
        ("Ataques", suite["attacks"]),
        ("Combinações", suite["comparisons_count"]),
        ("Cenários executados", suite["scenario_runs"]),
        ("Seed", configuration["seed"]),
        ("Intensidade", configuration["intensity"]),
        ("Usuários avaliados", configuration["sample_users"]),
        ("Top-k", configuration["top_k"]),
        ("Livro-alvo", configuration.get("target_book_id") or "Seleção automática"),
        ("Livros no snapshot", suite["dataset"]["books"]),
        ("Avaliações no snapshot", suite["dataset"]["ratings_snapshot"]),
        ("Tempo total (s)", suite["total_seconds"]),
        ("Pareamento", suite["method"]["pairing"]),
        (
            "Índice de tags vetorizado",
            "Sim" if suite["method"].get("vectorized_content_tags") else "Não",
        ),
        (
            "Baselines calculados",
            suite["method"].get("baseline_evaluations", "Não registrado"),
        ),
        (
            "Ataque bruto sem observabilidade",
            "Sim"
            if suite["method"].get(
                "raw_attack_retained_without_observability"
            )
            else "Não",
        ),
        ("Banco principal alterado", "Não"),
        (
            "Qualidade inicial",
            "Métricas do modelo antes da injeção do ataque, usando a mesma coorte e seed.",
        ),
        (
            "Sem observabilidade",
            "Métricas do lote atacado bruto, sem limpeza, detecção ou resposta automática.",
        ),
        (
            "Com observabilidade",
            "Métricas após o ataque com detecção e resposta automática quando aplicável.",
        ),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
    for row in sheet.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="334155")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"


def build_suite_workbook(suite: dict) -> BytesIO:
    """Cria o XLSX em memória para download autenticado."""

    workbook = Workbook()
    workbook.properties.title = "Comparação de qualidade dos recomendadores"
    workbook.properties.subject = "Qualidade inicial, sem e com observabilidade"
    workbook.properties.creator = "Recommendation API"
    _quality_sheet(workbook, suite)
    _operations_sheet(workbook, suite)
    _data_quality_sheet(workbook, suite)
    _method_sheet(workbook, suite)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
