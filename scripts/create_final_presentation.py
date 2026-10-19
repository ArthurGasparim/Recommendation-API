#!/usr/bin/env python3
"""Gera a apresentação final da iniciação científica a partir do modelo Unoeste."""

from __future__ import annotations

import argparse
from pathlib import Path
from statistics import mean

from PIL import Image
from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


EMU = 914400
GREEN = RGBColor(0x00, 0x66, 0x33)
DARK_GREEN = RGBColor(0x00, 0x4A, 0x2A)
LIME = RGBColor(0x91, 0xBD, 0x00)
PALE_GREEN = RGBColor(0xED, 0xF7, 0xE6)
BLUE = RGBColor(0x04, 0x42, 0x99)
PALE_BLUE = RGBColor(0xEA, 0xF3, 0xFA)
DARK = RGBColor(0x17, 0x24, 0x33)
TEXT = RGBColor(0x34, 0x45, 0x4F)
MUTED = RGBColor(0x70, 0x7D, 0x84)
LIGHT = RGBColor(0xF4, 0xF7, 0xF7)
BORDER = RGBColor(0xD9, 0xE3, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xD9, 0x3F, 0x3F)
ORANGE = RGBColor(0xF4, 0x87, 0x2D)
PURPLE = RGBColor(0x7B, 0x61, 0xFF)
PALE_ORANGE = RGBColor(0xFF, 0xF3, 0xE8)
PALE_PURPLE = RGBColor(0xF1, 0xEE, 0xFF)


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        relationship_id = slide_id.rId
        prs.part.drop_rel(relationship_id)
        slide_ids.remove(slide_id)


def set_shape_fill(shape, color: RGBColor, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_shape_line(shape, color: RGBColor, width: float = 1) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: RGBColor = TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font: str = "Aptos",
    margin: float = 0,
    italic: bool = False,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    # Impede que o LibreOffice expanda caixas longas horizontalmente.
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    paragraphs = []
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraphs.append(paragraph)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
    return box


def add_rich_text(slide, runs: list[dict], x, y, w, h, *, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for spec in runs:
        run = paragraph.add_run()
        run.text = spec["text"]
        run.font.name = "Aptos"
        run.font.size = Pt(spec.get("size", 18))
        run.font.bold = spec.get("bold", False)
        run.font.color.rgb = spec.get("color", TEXT)
    return box


def add_round_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_fill(shape, fill)
    set_shape_line(shape, line, 0.8)
    return shape


def add_circle(slide, x, y, diameter, fill, text, *, text_color=WHITE, size=18):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    set_shape_fill(shape, fill)
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos Display"
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = text_color
    return shape


def add_line(slide, x1, y1, x2, y2, color=BORDER, width=1, dash=None):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def add_slide_title(slide, section: str, title: str, subtitle: str | None = None):
    add_text(slide, section.upper(), 0.52, 0.25, 3.6, 0.22, size=9, color=LIME, bold=True)
    add_text(slide, title, 0.52, 0.48, 8.92, 0.45, size=21.5, color=DARK_GREEN, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.54, 0.91, 8.85, 0.28, size=11, color=MUTED)
    add_line(slide, 0.52, 1.19 if subtitle else 1.02, 9.46, 1.19 if subtitle else 1.02, BORDER, 0.8)


def add_section_pill(slide, text, x, y, w, *, fill=PALE_GREEN, color=GREEN):
    shape = add_round_rect(slide, x, y, w, 0.28, fill=fill, line=fill)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text.upper()
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = color
    return shape


def add_bullet_list(slide, items, x, y, w, h, *, size=16, color=TEXT, accent=LIME, gap=0.42):
    current_y = y
    for item in items:
        add_circle(slide, x, current_y + 0.08, 0.12, accent, "", size=1)
        add_text(slide, item, x + 0.23, current_y, w - 0.23, gap, size=size, color=color)
        current_y += gap


def add_picture_crop(slide, path: Path, x, y, w, h):
    with Image.open(path) as image:
        image_w, image_h = image.size
    frame_ratio = w / h
    image_ratio = image_w / image_h
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        crop = (1 - visible) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < frame_ratio:
        visible = image_ratio / frame_ratio
        crop = (1 - visible) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_note(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def add_source(slide, text: str) -> None:
    """Inclui a origem dos números sem competir com o conteúdo principal."""
    add_text(slide, text, 3.0, 5.34, 4.35, 0.14, size=5.8, color=MUTED, italic=True, align=PP_ALIGN.CENTER)


def load_measured_results(path: Path) -> dict:
    """Lê a bateria real e mantém os slides vinculados ao XLSX exportado."""
    workbook = load_workbook(path, data_only=True, read_only=True)

    def records(sheet_name: str) -> list[dict]:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(min_row=4, values_only=True)
        headers = next(rows)
        return [dict(zip(headers, row)) for row in rows if row[0] is not None]

    methodology = {
        row[0]: row[1]
        for row in workbook["Metodologia"].iter_rows(values_only=True)
        if row[0] not in (None, "Campo")
    }
    return {
        "quality": records("Qualidade"),
        "operation": records("Operação"),
        "data": records("Dados"),
        "methodology": methodology,
        "source": path.name,
    }


def measured_row(results: dict, sheet: str, model: str, attack: str) -> dict:
    for row in results[sheet]:
        if row["Modelo"] == model and row["Ataque"] == attack:
            return row
    raise KeyError(f"Resultado não encontrado: {sheet=} {model=} {attack=}")


def add_metric_triplet(
    slide,
    *,
    x: float,
    y: float,
    w: float,
    title: str,
    initial: str,
    attacked: str,
    observed: str,
    unit: str = "",
    colors=(BLUE, ORANGE, GREEN),
):
    add_text(slide, title, x, y, w, 0.2, size=9, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    labels = [("INICIAL", initial), ("SEM OBS.", attacked), ("COM OBS.", observed)]
    cell_w = (w - 0.12) / 3
    for index, ((label, value), color) in enumerate(zip(labels, colors)):
        cell_x = x + index * (cell_w + 0.06)
        add_round_rect(slide, cell_x, y + 0.28, cell_w, 0.68, fill=WHITE, line=color)
        add_text(slide, label, cell_x + 0.04, y + 0.39, cell_w - 0.08, 0.13, size=6.5, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, f"{value}{unit}", cell_x + 0.04, y + 0.61, cell_w - 0.08, 0.21, size=12, color=color, bold=True, align=PP_ALIGN.CENTER)


def add_explanation_card(slide, x, y, w, h, heading, body, *, color=GREEN, fill=WHITE):
    add_round_rect(slide, x, y, w, h, fill=fill, line=color)
    add_text(slide, heading.upper(), x + 0.16, y + 0.16, w - 0.32, 0.18, size=8, color=color, bold=True)
    add_text(slide, body, x + 0.16, y + 0.48, w - 0.32, h - 0.6, size=9.2, color=TEXT, line_spacing=1.05)


def add_model_result_table(
    slide,
    title: str,
    rows: list[tuple[str, str, str, str]],
    *,
    y: float = 2.56,
    first_header: str = "Modelo",
):
    """Tabela padronizada para comparar os três estados de cada modelo."""
    add_text(slide, title, 0.72, y, 3.4, 0.22, size=11, color=DARK_GREEN, bold=True)
    headers = [first_header, "Inicial", "Sem obs.", "Com obs."]
    widths = [3.35, 1.5, 1.5, 1.5]
    x_positions = [0.72, 4.07, 5.57, 7.07]
    header_y = y + 0.31
    for value, x, width in zip(headers, x_positions, widths):
        add_round_rect(slide, x, header_y, width, 0.38, fill=DARK_GREEN, line=DARK_GREEN, radius=False)
        add_text(slide, value, x + 0.05, header_y + 0.11, width - 0.1, 0.14, size=7.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    for index, row in enumerate(rows):
        row_y = header_y + 0.38 + index * 0.4
        for cell_index, (value, x, width) in enumerate(zip(row, x_positions, widths)):
            add_round_rect(slide, x, row_y, width, 0.38, fill=WHITE if index % 2 == 0 else LIGHT, line=BORDER, radius=False)
            add_text(
                slide,
                value,
                x + 0.05,
                row_y + 0.11,
                width - 0.1,
                0.14,
                size=7.5,
                color=DARK if cell_index == 0 else TEXT,
                bold=cell_index == 0,
                align=PP_ALIGN.CENTER,
            )


def make_cover(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    # Limpa os textos de exemplo antes de preencher o título. Placeholders
    # retornados pela biblioteca não preservam identidade de objeto.
    for shape in list(slide.shapes):
        if hasattr(shape, "text_frame"):
            shape.text_frame.clear()
    title.text = "Monitoramento de Qualidade dos Dados para\nPersonalização de Recomendação de Produtos"
    frame = title.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = "Aptos Display"
            run.font.size = Pt(31)
            run.font.bold = True
            run.font.color.rgb = WHITE

    add_text(slide, "Arthur Guarizi Gasparim", 2.5, 4.27, 5.0, 0.28, size=14, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Orientador: Prof. Dr. Mário Augusto Pazoti  •  2026", 1.8, 4.58, 6.4, 0.24, size=10, color=WHITE, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 20 segundos. Apresente o título, seu nome e o orientador. Antecipe que o trabalho integra recomendação, observabilidade e experimentação controlada.")


def slide_context(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Introdução", "Por que observar sistemas de recomendação?")
    add_rich_text(
        slide,
        [
            {"text": "A qualidade da recomendação depende diretamente da ", "size": 19, "color": DARK},
            {"text": "qualidade dos dados", "size": 19, "color": GREEN, "bold": True},
            {"text": " e da confiabilidade do sistema que os processa.", "size": 19, "color": DARK},
        ],
        0.65, 1.23, 8.7, 0.55,
    )
    cards = [
        ("01", "Dados incorretos", "Duplicatas, ausências e valores fora do domínio distorcem perfis e rankings.", ORANGE),
        ("02", "Manipulação", "Avaliações artificiais podem promover ou prejudicar itens e reduzir a confiança.", ORANGE),
        ("03", "Falhas distribuídas", "Latência e indisponibilidade dificultam localizar a causa apenas pela interface.", PURPLE),
    ]
    for index, (number, title, body, color) in enumerate(cards):
        x = 0.58 + index * 3.05
        add_round_rect(slide, x, 1.88, 2.75, 1.72, fill=WHITE, line=BORDER)
        add_circle(slide, x + 0.18, 2.08, 0.42, color, number, size=10)
        add_text(slide, title, x + 0.7, 2.05, 1.85, 0.28, size=15, color=DARK, bold=True)
        add_text(slide, body, x + 0.2, 2.58, 2.35, 0.76, size=11, color=MUTED, line_spacing=1.08)
    add_round_rect(slide, 0.72, 3.92, 8.56, 0.6, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "DEGRADAÇÃO SILENCIOSA", 0.95, 4.10, 1.75, 0.18, size=9, color=GREEN, bold=True)
    add_text(slide, "→", 2.65, 4.02, 0.4, 0.3, size=22, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "recomendações menos relevantes", 3.15, 4.08, 2.35, 0.24, size=12, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 5.48, 4.02, 0.4, 0.3, size=22, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "piora da experiência e da confiança", 5.88, 4.08, 2.9, 0.24, size=12, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 50 segundos. Explique que um recomendador pode continuar respondendo normalmente e, ainda assim, produzir listas piores. O problema científico é tornar essa degradação visível e acionável.")


def slide_objective(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Introdução", "Objetivo e pergunta de pesquisa")
    add_round_rect(slide, 0.72, 1.28, 8.55, 1.02, fill=PALE_GREEN, line=LIME)
    add_text(slide, "PERGUNTA DE PESQUISA", 0.98, 1.47, 1.55, 0.2, size=9, color=GREEN, bold=True)
    add_text(
        slide,
        "Como a observabilidade e o monitoramento da qualidade dos dados auxiliam na detecção de falhas e preservam a qualidade de sistemas de recomendação?",
        2.45, 1.4, 6.35, 0.75, size=14, color=DARK_GREEN, bold=True,
    )
    goals = [
        ("RECOMENDAR", "Construir um sistema funcional com conteúdo e filtragem colaborativa.", GREEN),
        ("OBSERVAR", "Coletar sinais de aplicação, dados e infraestrutura para explicar degradações.", BLUE),
        ("COMPARAR", "Simular falhas no tempo e contrastar operação sem defesa e com resposta observável.", PURPLE),
    ]
    for index, (label, body, color) in enumerate(goals):
        x = 0.72 + index * 2.91
        add_round_rect(slide, x, 2.62, 2.63, 1.45, fill=WHITE, line=BORDER)
        add_section_pill(slide, label, x + 0.18, 2.83, 1.16, fill=PALE_BLUE if color == BLUE else PALE_GREEN, color=color)
        add_text(slide, body, x + 0.18, 3.26, 2.25, 0.55, size=11, color=TEXT)
    add_text(slide, "Contribuição esperada", 0.78, 4.28, 1.55, 0.24, size=11, color=GREEN, bold=True)
    add_text(slide, "Um ambiente reproduzível para estudar qualidade, resiliência e custo da observabilidade em recomendadores.", 2.25, 4.25, 6.55, 0.32, size=13, color=DARK, bold=True)
    add_note(slide, "Tempo sugerido: 50 segundos. Leia a pergunta de pesquisa e resuma os três verbos do trabalho: recomendar, observar e comparar. Reforce que a contribuição é tanto técnica quanto experimental.")


def slide_observability(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Observabilidade", "Do sinal à recuperação da qualidade")
    pillars = [
        ("M", "Métricas", "Tendências, SLOs, latência, erros e qualidade dos dados.", GREEN),
        ("L", "Logs", "Contexto detalhado de eventos, versões e causas prováveis.", BLUE),
        ("T", "Tracing", "Caminho da requisição entre API, banco e inferência.", PURPLE),
    ]
    for i, (letter, title, body, color) in enumerate(pillars):
        x = 0.72 + i * 2.95
        add_round_rect(slide, x, 1.35, 2.62, 1.55, fill=WHITE, line=BORDER)
        add_circle(slide, x + 0.18, 1.58, 0.5, color, letter, size=17)
        add_text(slide, title, x + 0.82, 1.58, 1.45, 0.25, size=16, color=DARK, bold=True)
        add_text(slide, body, x + 0.2, 2.14, 2.22, 0.55, size=10.5, color=MUTED)
    steps = [
        ("SINAL", "mudança medida"),
        ("DETECÇÃO", "alerta"),
        ("DIAGNÓSTICO", "causa"),
        ("AÇÃO", "bloqueio / rollback"),
        ("RECUPERAÇÃO", "qualidade restaurada"),
    ]
    for i, (label, detail) in enumerate(steps):
        x = 0.54 + i * 1.84
        add_round_rect(slide, x, 3.35, 1.58, 0.78, fill=PALE_GREEN if i in (0, 4) else LIGHT, line=BORDER)
        add_text(slide, label, x + 0.08, 3.5, 1.42, 0.2, size=9, color=GREEN if i in (0, 4) else DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.08, 3.78, 1.42, 0.18, size=8, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 4:
            add_text(slide, "→", x + 1.58, 3.55, 0.26, 0.24, size=17, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Observabilidade não melhora o algoritmo sozinha; ela reduz o tempo de exposição ao problema ao orientar a resposta.", 0.95, 4.34, 8.1, 0.3, size=12, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 55 segundos. Defina os três pilares e percorra o fluxo da esquerda para a direita. Destaque a distinção central: instrumentação é passiva; o ganho de qualidade ocorre pela ação corretiva orientada pelos sinais.")


def slide_recommendation(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Recomendação", "Abordagens implementadas e modelos comparados")
    # Conteúdo
    add_round_rect(slide, 0.58, 1.32, 4.25, 1.62, fill=PALE_BLUE, line=RGBColor(0xB9, 0xD7, 0xEA))
    add_section_pill(slide, "Baseado em conteúdo", 0.83, 1.57, 1.65, fill=WHITE, color=BLUE)
    add_text(slide, "Tags + atributos do livro", 0.83, 2.02, 3.55, 0.25, size=14, color=DARK, bold=True)
    add_text(slide, "Compara características dos itens e o histórico do usuário; inclui tratamento de cold start.", 0.83, 2.37, 3.55, 0.4, size=10.5, color=MUTED)
    # Colaborativo
    add_round_rect(slide, 5.17, 1.32, 4.25, 1.62, fill=PALE_GREEN, line=RGBColor(0xC8, 0xE2, 0xB2))
    add_section_pill(slide, "Filtragem colaborativa", 5.42, 1.57, 1.78, fill=WHITE, color=GREEN)
    add_text(slide, "Similaridade item-item", 5.42, 2.02, 3.55, 0.25, size=14, color=DARK, bold=True)
    add_text(slide, "Usa uma matriz esparsa usuário-item e retorna IDs reais, excluindo itens já avaliados.", 5.42, 2.37, 3.55, 0.4, size=10.5, color=MUTED)

    add_text(slide, "Perfis previstos no protocolo comparativo", 0.65, 3.23, 4.3, 0.25, size=11, color=DARK_GREEN, bold=True)
    models = [
        ("Popularidade", "baseline", MUTED),
        ("TF-IDF", "conteúdo", BLUE),
        ("Item-item", "colaborativo", PURPLE),
        ("ALS", "fatores latentes", ORANGE),
        ("Híbrido", "combinação", GREEN),
    ]
    for i, (name, detail, color) in enumerate(models):
        x = 0.6 + i * 1.84
        add_round_rect(slide, x, 3.63, 1.62, 0.72, fill=WHITE, line=BORDER)
        add_circle(slide, x + 0.12, 3.83, 0.17, color, "", size=1)
        add_text(slide, name, x + 0.38, 3.72, 1.1, 0.2, size=10, color=DARK, bold=True)
        add_text(slide, detail, x + 0.38, 4.0, 1.1, 0.18, size=8, color=MUTED)
    add_text(slide, "Nota: ALS e híbrido aparecem no simulador como perfis exploratórios; a comparação científica depende de treinamento e avaliação medidos.", 0.7, 4.48, 8.55, 0.25, size=9.5, color=RED, italic=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 60 segundos. Explique as duas abordagens funcionais do protótipo. Em seguida, apresente os cinco perfis do protocolo e seja explícito: ALS e híbrido ainda não constituem resultados científicos medidos; no laboratório atual são perfis exploratórios.")


def slide_experiment(prs, lab_top: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Observabilidade + Recomendação", "Desenho experimental: ataques no tempo")
    add_round_rect(slide, 0.55, 1.26, 5.54, 3.24, fill=WHITE, line=BORDER)
    add_picture_crop(slide, lab_top, 0.58, 1.29, 5.48, 3.18)
    right_x = 6.35
    add_text(slide, "Cenários controlados", right_x, 1.34, 2.7, 0.28, size=15, color=DARK_GREEN, bold=True)
    attacks = ["Push e nuke attack", "Dados inválidos e duplicados", "Drift de distribuição", "Latência e indisponibilidade"]
    add_bullet_list(slide, attacks, right_x, 1.76, 2.8, 1.65, size=10.5, gap=0.38)
    add_round_rect(slide, right_x, 3.39, 2.88, 0.87, fill=PALE_GREEN, line=LIME)
    add_text(slide, "MESMOS DADOS • MESMA SEED", right_x + 0.16, 3.56, 2.56, 0.18, size=8, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "sem defesa  ×  resposta observável", right_x + 0.14, 3.83, 2.6, 0.2, size=10.5, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "A simulação valida o protocolo e a visualização; os valores atuais são sintéticos e determinísticos.", 6.38, 4.39, 2.8, 0.32, size=8.5, color=RED, italic=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 60 segundos. Mostre que cada ataque tem início, duração e intensidade. A comparação é pareada: mesmas condições, variando apenas a resposta orientada por observabilidade. Reforce que esta tela valida o desenho experimental, não substitui medições reais.")


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Implementação", "Arquitetura do protótipo")
    # Nuvem / cluster
    cluster = add_round_rect(slide, 0.5, 1.24, 8.98, 3.2, fill=LIGHT, line=BORDER)
    cluster.line.dash_style = 1
    add_text(slide, "AMBIENTE LOCAL / KUBERNETES / EKS", 0.74, 1.38, 3.1, 0.2, size=8.5, color=GREEN, bold=True)

    nodes = [
        (0.78, 2.08, 1.3, 0.72, "Usuário", "browser", BLUE),
        (2.36, 2.08, 1.55, 0.72, "React + Nginx", "frontend", GREEN),
        (4.2, 2.08, 1.55, 0.72, "FastAPI", "API + auth", DARK_GREEN),
        (6.04, 2.08, 1.55, 0.72, "PostgreSQL", "persistência", PURPLE),
    ]
    for x, y, w, h, title, detail, color in nodes:
        add_round_rect(slide, x, y, w, h, fill=WHITE, line=color)
        add_text(slide, title, x + 0.08, y + 0.13, w - 0.16, 0.2, size=11.5, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.08, y + 0.42, w - 0.16, 0.15, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    for x in [2.08, 3.91, 5.75]:
        add_text(slide, "→", x, 2.25, 0.28, 0.24, size=17, color=LIME, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 4.2, 3.22, 1.55, 0.62, fill=PALE_GREEN, line=LIME)
    add_text(slide, "Recomendadores", 4.3, 3.37, 1.35, 0.18, size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "↕", 4.82, 2.84, 0.3, 0.3, size=18, color=LIME, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 7.95, 1.62, 1.23, 2.28, fill=WHITE, line=BORDER)
    add_text(slide, "OBSERVABILIDADE", 8.04, 1.78, 1.04, 0.2, size=6.8, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    for i, (name, color) in enumerate([("Prometheus", GREEN), ("Grafana", ORANGE), ("Alertas", PURPLE)]):
        add_circle(slide, 8.15, 2.25 + i * 0.48, 0.16, color, "", size=1)
        add_text(slide, name, 8.38, 2.19 + i * 0.48, 0.7, 0.2, size=7.8, color=TEXT)
    add_line(slide, 5.75, 1.93, 7.95, 1.93, GREEN, 1.5)
    add_text(slide, "/metrics", 6.55, 1.72, 0.7, 0.18, size=8, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    add_text(slide, "Controles incorporados", 0.72, 4.58, 1.55, 0.2, size=9, color=GREEN, bold=True)
    add_text(slide, "tokens assinados  •  autorização administrativa  •  validação  •  revisão do modelo  •  PVC + backup", 2.15, 4.55, 6.8, 0.22, size=9.5, color=TEXT)
    add_note(slide, "Tempo sugerido: 60 segundos. Percorra a requisição do usuário ao banco. Mostre que o recomendador é reconstruído por revisão e que Prometheus coleta métricas da API. Cite execução local com Docker e implantação reproduzível em Kubernetes/EKS.")


def slide_data(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Implementação", "Dados e pipeline do sistema", "Base Goodbooks normalizada e carregada no banco local")
    stats = [
        ("10.000", "livros", GREEN),
        ("979.478", "avaliações", BLUE),
        ("34.252", "tags", PURPLE),
        ("999.904", "relações livro–tag", ORANGE),
        ("912.705", "itens “quero ler”", DARK_GREEN),
    ]
    for i, (value, label, color) in enumerate(stats):
        x = 0.48 + i * 1.86
        add_round_rect(slide, x, 1.42, 1.66, 0.86, fill=WHITE, line=BORDER)
        add_text(slide, value, x + 0.1, 1.59, 1.46, 0.25, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.08, 1.96, 1.5, 0.16, size=8, color=MUTED, align=PP_ALIGN.CENTER)

    steps = [
        ("1", "CSV", "entrada"),
        ("2", "Validação", "tipos e domínio"),
        ("3", "Consolidar", "duplicatas"),
        ("4", "Banco", "transação"),
        ("5", "Snapshot", "revisão"),
        ("6", "Ranking", "Top-K"),
    ]
    for i, (number, label, detail) in enumerate(steps):
        x = 0.56 + i * 1.5
        add_circle(slide, x, 2.83, 0.4, GREEN if i in (0, 5) else LIME, number, size=10)
        add_text(slide, label, x - 0.28, 3.34, 0.98, 0.2, size=9, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x - 0.22, 3.63, 0.86, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            add_line(slide, x + 0.43, 3.03, x + 1.42, 3.03, BORDER, 1.3)
    add_round_rect(slide, 0.7, 4.16, 8.6, 0.48, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "Correção aplicada: chaves compostas repetidas são consolidadas antes da inserção, preservando as restrições do banco.", 0.93, 4.3, 8.15, 0.2, size=10, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 50 segundos. Apresente o volume efetivamente carregado após consolidação. Explique rapidamente o pipeline e cite que a descoberta de duplicatas no dataset levou a uma correção validada por teste de regressão.")


def slide_system(prs, catalog: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Implementação", "Sistema web funcional")
    add_round_rect(slide, 0.48, 1.17, 6.68, 3.55, fill=WHITE, line=BORDER)
    add_picture_crop(slide, catalog, 0.55, 1.24, 6.54, 3.41)
    add_text(slide, "Fluxos disponíveis", 7.45, 1.36, 1.8, 0.28, size=15, color=DARK_GREEN, bold=True)
    add_bullet_list(
        slide,
        [
            "Autenticação com senha e token",
            "Catálogo e consulta de livros",
            "Avaliações, tags e lista de leitura",
            "Recomendação por item e por usuário",
            "Área administrativa protegida",
        ],
        7.43, 1.83, 1.95, 2.2, size=9.5, gap=0.4,
    )
    add_round_rect(slide, 7.38, 4.05, 1.94, 0.52, fill=PALE_GREEN, line=LIME)
    add_text(slide, "React + FastAPI integrados", 7.53, 4.21, 1.64, 0.18, size=8.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 45 segundos. Mostre a tela real e destaque que o usuário pode navegar pelo catálogo, avaliar livros e solicitar recomendações. Cite que autorização administrativa saiu do frontend e passou a ser validada no backend.")


def slide_lab(prs, lab_graph: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Implementação", "Laboratório de Resiliência", "Visualização temporal de ataques, detecção e recuperação")
    add_round_rect(slide, 0.48, 1.26, 9.02, 3.48, fill=WHITE, line=BORDER)
    add_picture_crop(slide, lab_graph, 0.55, 1.33, 8.88, 3.34)
    add_section_pill(slide, "ataque", 6.68, 1.48, 0.70, fill=RGBColor(0xFE, 0xED, 0xE8), color=RED)
    add_section_pill(slide, "detecção", 7.45, 1.48, 0.82, fill=PALE_GREEN, color=GREEN)
    add_section_pill(slide, "recuperação", 8.34, 1.48, 1.10, fill=PALE_BLUE, color=BLUE)
    add_note(slide, "Tempo sugerido: 60 segundos. Explique a linha vermelha pontilhada sem defesa e a linha verde com resposta observável. Mostre as faixas dos ataques e os marcadores de detecção. Ressalte novamente que o gráfico atual é exploratório; sua função é preparar a execução dos experimentos reais.")


def slide_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Validação funcional e técnica")
    top_cards = [
        ("15/15", "testes backend", GREEN),
        ("2/2", "testes frontend", BLUE),
        ("10.000", "livros processados", PURPLE),
        ("2", "bancos validados", ORANGE),
    ]
    for i, (value, label, color) in enumerate(top_cards):
        x = 0.63 + i * 2.27
        add_round_rect(slide, x, 1.33, 2.03, 0.94, fill=WHITE, line=BORDER)
        add_text(slide, value, x + 0.12, 1.53, 1.79, 0.27, size=20, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.1, 1.93, 1.83, 0.16, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)

    columns = [
        (0.67, "RECOMENDAÇÃO", ["IDs reais preservados", "Cold start tratado", "Itens já vistos excluídos", "Modelo atualizado por revisão"]),
        (3.54, "SEGURANÇA", ["Senha com PBKDF2 + salt", "Token assinado e expiração", "Autorização administrativa", "Entradas validadas"]),
        (6.41, "REPRODUTIBILIDADE", ["Docker e build aprovados", "SQLite e PostgreSQL 16", "Kustomize renderizado", "PVC, backup e restauração"]),
    ]
    for x, heading, bullets in columns:
        add_round_rect(slide, x, 2.65, 2.5, 1.66, fill=LIGHT, line=BORDER)
        add_text(slide, heading, x + 0.16, 2.85, 2.18, 0.2, size=9, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        add_bullet_list(slide, bullets, x + 0.2, 3.22, 2.1, 0.9, size=8.7, gap=0.25)

    add_text(slide, "Exemplo validado", 0.72, 4.47, 1.2, 0.18, size=9, color=GREEN, bold=True)
    add_text(slide, "The Hunger Games → Catching Fire, Mockingjay e The Hunger Games Box Set", 1.83, 4.44, 7.15, 0.22, size=10.5, color=DARK, bold=True)
    add_note(slide, "Tempo sugerido: 60 segundos. Estes são resultados efetivamente validados. Destaque testes, integração com dois bancos e os mecanismos de segurança. Dê o exemplo de recomendação de The Hunger Games para tornar o resultado concreto.")


def slide_simulation_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "O que o laboratório já permite analisar")
    items = [
        ("5", "perfis de modelos", "comparação simultânea", GREEN),
        ("6", "tipos de ataque", "com sobreposição temporal", RED),
        ("7", "métricas", "qualidade + operação", BLUE),
        ("2", "condições", "sem defesa × observável", PURPLE),
    ]
    for i, (value, label, detail, color) in enumerate(items):
        x = 0.68 + i * 2.25
        add_circle(slide, x, 1.46, 0.62, color, value, size=19)
        add_text(slide, label, x + 0.75, 1.47, 1.32, 0.4, size=9.5, color=DARK, bold=True)
        add_text(slide, detail, x + 0.75, 1.96, 1.32, 0.27, size=7.3, color=MUTED)

    add_round_rect(slide, 0.7, 2.55, 4.05, 1.47, fill=PALE_GREEN, line=LIME)
    add_text(slide, "HIPÓTESE VISUALIZADA", 0.95, 2.82, 1.55, 0.18, size=8.5, color=GREEN, bold=True)
    add_text(slide, "A queda inicial pode ser semelhante, mas a resposta orientada pelos sinais reduz o tempo abaixo do SLO e a perda acumulada de qualidade.", 0.95, 3.17, 3.54, 0.52, size=12, color=DARK_GREEN, bold=True)

    add_round_rect(slide, 5.15, 2.55, 4.12, 1.47, fill=RGBColor(0xFF, 0xF4, 0xE8), line=ORANGE)
    add_text(slide, "LIMITE CIENTÍFICO", 5.4, 2.82, 1.4, 0.18, size=8.5, color=ORANGE, bold=True)
    add_text(slide, "As curvas atuais são sintéticas. Elas comprovam o funcionamento do simulador, não a superioridade real de um modelo ou defesa.", 5.4, 3.17, 3.62, 0.52, size=12, color=DARK, bold=True)

    add_text(slide, "Métricas preparadas", 0.75, 4.31, 1.7, 0.2, size=10, color=GREEN, bold=True)
    add_text(slide, "NDCG@10 • Precision@10 • Recall@10 • cobertura • latência p95 • disponibilidade • qualidade dos dados", 2.35, 4.28, 6.75, 0.32, size=9.3, color=TEXT, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 55 segundos. Separe claramente implementação e evidência científica. O laboratório já executa a comparação e calcula perda acumulada, MTTD e overhead. Porém, as curvas ainda são sintéticas; a próxima etapa é alimentar a mesma visualização com medições reais.")


def slide_conclusion(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Conclusões e continuidade")
    contributions = [
        ("1", "Protótipo integrado", "Recomendação, interface, segurança, banco e nuvem em uma base reproduzível."),
        ("2", "Observabilidade acionável", "Métricas e alertas ligados a um protocolo de detecção, diagnóstico e recuperação."),
        ("3", "Laboratório experimental", "Ataques programáveis e comparação pareada entre modelos e condições."),
    ]
    for i, (number, title, body) in enumerate(contributions):
        y = 1.35 + i * 1.02
        add_circle(slide, 0.72, y, 0.48, GREEN, number, size=13)
        add_text(slide, title, 1.42, y + 0.01, 3.7, 0.24, size=13.5, color=DARK_GREEN, bold=True)
        add_text(slide, body, 1.42, y + 0.36, 3.65, 0.34, size=10.5, color=MUTED)

    add_round_rect(slide, 5.45, 1.34, 3.8, 2.94, fill=LIGHT, line=BORDER)
    add_text(slide, "PRÓXIMAS MEDIÇÕES", 5.75, 1.65, 2.0, 0.2, size=9, color=GREEN, bold=True)
    add_bullet_list(
        slide,
        [
            "Treinar ALS e híbrido",
            "Substituir curvas sintéticas por medições",
            "Executar ataques em cópias controladas",
            "Completar logs e tracing",
            "Repetir o protocolo no EKS",
        ],
        5.72, 2.05, 3.05, 1.7, size=10, gap=0.34,
    )
    add_round_rect(slide, 0.78, 4.46, 8.42, 0.38, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "Síntese: observar não evita a falha inicial, mas torna possível detectá-la, explicá-la e recuperar o sistema mais cedo.", 0.98, 4.56, 8.02, 0.18, size=10.5, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 55 segundos. Feche com as três contribuições concretas. Depois, reconheça as próximas medições necessárias. A frase final resume a tese do trabalho: observabilidade reduz o tempo de exposição à degradação.")


def slide_metrics_explained(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Métricas", "Como a qualidade e a operação são medidas")
    metrics = [
        ("NDCG@10", "Relevância do Top-10, dando mais peso aos acertos nas primeiras posições.", BLUE),
        ("Precision@10", "Parcela dos 10 itens recomendados que é relevante para o usuário.", GREEN),
        ("Recall@10", "Parcela dos itens relevantes que aparece entre as 10 recomendações.", PURPLE),
        ("Hit Rate@10", "Percentual de usuários que recebe ao menos um item relevante no Top-10.", ORANGE),
        ("Cobertura", "Fração do catálogo que aparece nas recomendações; indica diversidade de exposição.", DARK_GREEN),
        ("Latência p95", "Tempo abaixo do qual terminam 95% das requisições; evidencia a cauda lenta.", BLUE),
        ("Disponibilidade", "Percentual de requisições concluídas com sucesso durante o cenário.", GREEN),
        ("Qualidade dos dados", "Score composto por domínio válido, integridade referencial e unicidade.", PURPLE),
    ]
    for index, (name, definition, color) in enumerate(metrics):
        row, column = divmod(index, 4)
        x = 0.48 + column * 2.38
        y = 1.28 + row * 1.54
        add_round_rect(slide, x, y, 2.18, 1.28, fill=WHITE, line=color)
        add_text(slide, name, x + 0.16, y + 0.17, 1.86, 0.22, size=12, color=color, bold=True)
        add_text(slide, definition, x + 0.16, y + 0.51, 1.86, 0.53, size=8.6, color=TEXT, line_spacing=1.02)
    add_round_rect(slide, 0.72, 4.47, 8.56, 0.43, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "Princípio usado: cada ataque é avaliado pela métrica que ele realmente pretende degradar.", 0.95, 4.59, 8.1, 0.18, size=10.5, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 65 segundos. Não leia todas as definições. Explique NDCG, p95 e qualidade dos dados e use as demais como referência. O ponto central é que relevância, dados e operação são dimensões diferentes.")


def slide_recommendation_measured(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Recomendação", "Três modelos executados na bateria real")
    attack = "Push attack"
    cards = [
        (
            "Popularidade regularizada",
            "Baseline que ordena itens pela popularidade, com regularização para reduzir o efeito de poucos votos.",
            "Lista global; não personaliza por perfil.",
            measured_row(results, "quality", "Popularidade regularizada", attack),
            ORANGE,
        ),
        (
            "Conteúdo por tags",
            "Representa tags por TF-IDF e usa similaridade de cosseno com o histórico do usuário.",
            "Responde às características dos livros.",
            measured_row(results, "quality", "Conteúdo por tags", attack),
            BLUE,
        ),
        (
            "Colaborativo item-item",
            "Compara colunas da matriz usuário–item para encontrar livros avaliados de forma semelhante.",
            "Responde ao comportamento coletivo.",
            measured_row(results, "quality", "Colaborativo item-item", attack),
            PURPLE,
        ),
    ]
    for index, (name, explanation, behavior, row, color) in enumerate(cards):
        x = 0.55 + index * 3.08
        add_round_rect(slide, x, 1.32, 2.77, 2.86, fill=WHITE, line=color)
        add_circle(slide, x + 0.2, 1.56, 0.42, color, str(index + 1), size=11)
        add_text(slide, name, x + 0.76, 1.53, 1.78, 0.5, size=13, color=DARK, bold=True)
        add_text(slide, explanation, x + 0.2, 2.24, 2.37, 0.75, size=9.2, color=TEXT)
        add_text(slide, behavior, x + 0.2, 3.05, 2.37, 0.32, size=8.5, color=MUTED, italic=True)
        add_text(slide, f"NDCG inicial  {row['NDCG - Inicial']:.6f}", x + 0.2, 3.55, 2.35, 0.2, size=10, color=color, bold=True)
        add_text(slide, f"Cobertura inicial  {row['Cobertura - Inicial']:.4f}", x + 0.2, 3.84, 2.35, 0.18, size=8.5, color=MUTED)
    add_text(slide, "TF-IDF: destaca tags características.  •  Cosseno: mede a proximidade entre vetores.  •  Regularização: reduz extremos com pouca evidência.", 0.72, 4.54, 8.55, 0.34, size=9.2, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: bateria experimental medida em 19/10/2026 — condição inicial, Top-10.")
    add_note(slide, "Tempo sugerido: 65 segundos. Explique cada técnica em uma frase. Mostre que conteúdo teve o maior NDCG inicial, item-item a maior cobertura e popularidade funcionou como baseline e fallback.")


def slide_real_protocol(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    meta = results["methodology"]
    add_slide_title(slide, "Metodologia", "Bateria pareada e medida", "Mesma seed, mesma coorte e mesmo snapshot nos dois modos")
    stats = [
        (str(meta["Modelos"]), "modelos", BLUE),
        (str(meta["Ataques"]), "ataques", ORANGE),
        (str(meta["Combinações"]), "comparações", PURPLE),
        (str(meta["Cenários executados"]), "execuções", GREEN),
        (f"{meta['Tempo total (s)']:.1f}s", "tempo total", DARK_GREEN),
    ]
    for index, (value, label, color) in enumerate(stats):
        x = 0.48 + index * 1.88
        add_round_rect(slide, x, 1.34, 1.66, 0.8, fill=WHITE, line=color)
        add_text(slide, value, x + 0.08, 1.49, 1.5, 0.25, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.08, 1.85, 1.5, 0.16, size=8, color=MUTED, align=PP_ALIGN.CENTER)

    stages = [
        ("QUALIDADE INICIAL", "modelo antes da injeção", BLUE),
        ("SEM OBSERVABILIDADE", "lote atacado permanece bruto", ORANGE),
        ("COM OBSERVABILIDADE", "detecção + resposta automática", GREEN),
    ]
    for index, (label, detail, color) in enumerate(stages):
        x = 0.75 + index * 3.06
        add_round_rect(slide, x, 2.62, 2.62, 0.96, fill=WHITE, line=color)
        add_text(slide, label, x + 0.12, 2.83, 2.38, 0.18, size=8.5, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.12, 3.17, 2.38, 0.18, size=8, color=MUTED, align=PP_ALIGN.CENTER)
        if index < 2:
            add_text(slide, "→", x + 2.64, 2.88, 0.36, 0.28, size=20, color=LIME, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 0.75, 3.88, 8.5, 0.78, fill=PALE_GREEN, line=LIME)
    add_text(slide, f"Seed {meta['Seed']}  •  intensidade {meta['Intensidade']:.0%}  •  {meta['Usuários avaliados']} usuários  •  Top-{meta['Top-k']}  •  alvo automático", 0.95, 4.05, 8.1, 0.2, size=10.5, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "O ataque ocorre em uma cópia controlada; o banco principal não é alterado.", 1.2, 4.35, 7.6, 0.18, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: aba Metodologia do XLSX exportado pelo sistema em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 60 segundos. Mostre que a comparação é pareada. O banco real origina um snapshot, mas cada ataque atua somente sobre uma cópia controlada. Isso torna o experimento repetível e seguro.")


def slide_attack_map(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Metodologia", "Ataques, sinais, respostas e métricas")
    attacks = [
        ("Push", "promove item-alvo com perfis artificiais", "perfis atípicos → rollback", "exposição/rank do alvo", ORANGE),
        ("Nuke", "rebaixa item-alvo com notas artificiais", "perfis atípicos → rollback", "exposição/rank do alvo", PURPLE),
        ("Corrupção", "injeta inválidos, órfãos e duplicatas", "regras de dados → quarentena", "score de qualidade", BLUE),
        ("Drift", "desloca a distribuição das notas", "mudança estatística → rollback", "média + NDCG", GREEN),
        ("Latência", "adiciona atraso ao caminho crítico", "p95 alto → circuit breaker", "p95 e p99", DARK_GREEN),
        ("Indisponibilidade", "simula falha de dependência", "erros → fallback", "disponibilidade + falhas", ORANGE),
    ]
    for index, (name, injection, response, metric, color) in enumerate(attacks):
        row, column = divmod(index, 3)
        x = 0.52 + column * 3.08
        y = 1.27 + row * 1.65
        add_round_rect(slide, x, y, 2.82, 1.43, fill=WHITE, line=color)
        add_text(slide, name, x + 0.16, y + 0.16, 1.16, 0.22, size=12, color=color, bold=True)
        add_text(slide, injection, x + 0.16, y + 0.48, 2.48, 0.3, size=8.3, color=TEXT)
        add_text(slide, response, x + 0.16, y + 0.84, 2.48, 0.22, size=7.7, color=DARK_GREEN, bold=True)
        add_text(slide, f"Métrica: {metric}", x + 0.16, y + 1.15, 2.48, 0.16, size=7.5, color=MUTED, italic=True)
    add_round_rect(slide, 0.72, 4.72, 8.56, 0.32, fill=PALE_BLUE, line=PALE_BLUE)
    add_text(slide, "Observabilidade só demonstra benefício quando sinal, resposta e métrica estão alinhados ao efeito do ataque.", 0.95, 4.8, 8.1, 0.16, size=9.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 65 segundos. Percorra a grade sem detalhar os resultados; os próximos slides fazem isso. Reforce que ataques de infraestrutura não precisam reduzir NDCG, e ataques ao alvo exigem métricas específicas de exposição.")


def slide_push_nuke_results(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Push e nuke attack: manipulação do item-alvo")
    data_push = measured_row(results, "data", "Popularidade regularizada", "Push attack")
    data_nuke = measured_row(results, "data", "Popularidade regularizada", "Nuke attack")
    add_explanation_card(slide, 0.55, 1.26, 4.25, 0.96, "O ataque", "Perfis artificiais tentam promover (push) ou rebaixar (nuke) um livro-alvo por meio de avaliações coordenadas.", color=ORANGE, fill=PALE_ORANGE)
    add_explanation_card(slide, 5.08, 1.26, 4.35, 0.96, "A resposta observável", "O aumento de perfis atípicos é detectado e o lote atacado sofre rollback para o snapshot íntegro.", color=GREEN, fill=PALE_GREEN)

    add_metric_triplet(
        slide, x=0.65, y=2.53, w=4.1, title="Push — nota média do snapshot",
        initial=f"{data_push['Nota média - Inicial']:.3f}", attacked=f"{data_push['Nota média - Sem observabilidade']:.3f}", observed=f"{data_push['Nota média - Com observabilidade']:.3f}",
    )
    add_metric_triplet(
        slide, x=5.25, y=2.53, w=4.1, title="Nuke — nota média do snapshot",
        initial=f"{data_nuke['Nota média - Inicial']:.3f}", attacked=f"{data_nuke['Nota média - Sem observabilidade']:.3f}", observed=f"{data_nuke['Nota média - Com observabilidade']:.3f}",
    )
    add_round_rect(slide, 0.72, 3.92, 8.56, 0.85, fill=PALE_BLUE, line=BLUE)
    add_text(slide, "RESULTADO E LIMITE DA MÉTRICA", 0.95, 4.09, 2.05, 0.18, size=8, color=BLUE, bold=True)
    add_text(slide, "O NDCG@10 permaneceu igual nos três modelos. Isso não invalida o ataque: o NDCG mede o item relevante retido, e não a promoção ou supressão do alvo. Para esses ataques, deve-se acrescentar exposição, posição e taxa de sucesso do alvo.", 2.8, 4.03, 6.1, 0.54, size=9.3, color=DARK, bold=True)
    add_source(slide, "Fonte: abas Dados, Qualidade e Operação da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 70 segundos. Explique o ganho artificial na média e o rollback. Seja explícito: NDCG inalterado não significa defesa perfeita; significa que a métrica atual não observa diretamente o objetivo do atacante. Essa é uma conclusão metodológica do experimento.")


def slide_corruption_results(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Corrupção de dados: a qualidade cai e é restaurada")
    data = measured_row(results, "data", "Conteúdo por tags", "Corrupção de dados")
    quality = measured_row(results, "quality", "Conteúdo por tags", "Corrupção de dados")
    add_metric_triplet(
        slide, x=0.6, y=1.3, w=4.25, title="Score de qualidade dos dados",
        initial=f"{data['Score de qualidade - Inicial']:.3f}", attacked=f"{data['Score de qualidade - Sem observabilidade']:.3f}", observed=f"{data['Score de qualidade - Com observabilidade']:.3f}",
    )
    add_metric_triplet(
        slide, x=5.15, y=1.3, w=4.25, title="Conteúdo — NDCG@10",
        initial=f"{quality['NDCG - Inicial']:.6f}", attacked=f"{quality['NDCG - Sem observabilidade']:.6f}", observed=f"{quality['NDCG - Com observabilidade']:.6f}",
    )
    rates = [
        ("7,69%", "domínio inválido"),
        ("7,69%", "referências órfãs"),
        ("37,5%", "duplicidade"),
        ("−9,5%", "NDCG de conteúdo"),
    ]
    for index, (value, label) in enumerate(rates):
        x = 0.6 + index * 2.25
        add_round_rect(slide, x, 2.8, 2.02, 0.72, fill=PALE_ORANGE if index < 3 else PALE_PURPLE, line=ORANGE if index < 3 else PURPLE)
        add_text(slide, value, x + 0.1, 2.95, 0.72, 0.22, size=14, color=ORANGE if index < 3 else PURPLE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.82, 2.94, 1.08, 0.3, size=7.8, color=TEXT, bold=True)
    add_explanation_card(slide, 0.65, 3.82, 4.15, 0.92, "Como a observabilidade ajuda", "Regras de dados detectam o lote; a quarentena impede sua incorporação e restaura o score 1,000.", color=GREEN, fill=PALE_GREEN)
    add_explanation_card(slide, 5.12, 3.82, 4.2, 0.92, "Leitura correta", "A oscilação de uma métrica de ranking isolada não substitui a medição direta da saúde dos dados.", color=BLUE, fill=PALE_BLUE)
    add_source(slide, "Fonte: abas Dados e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 65 segundos. Mostre a queda do score de qualidade e a perda de NDCG no modelo de conteúdo. A resposta é a quarentena, não uma alteração no algoritmo de recomendação.")


def slide_drift_results(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Drift: a distribuição muda e o ranking sente")
    data = measured_row(results, "data", "Conteúdo por tags", "Drift de distribuição")
    content = measured_row(results, "quality", "Conteúdo por tags", "Drift de distribuição")
    item = measured_row(results, "quality", "Colaborativo item-item", "Drift de distribuição")
    add_metric_triplet(
        slide, x=0.58, y=1.29, w=4.25, title="Nota média do snapshot",
        initial=f"{data['Nota média - Inicial']:.3f}", attacked=f"{data['Nota média - Sem observabilidade']:.3f}", observed=f"{data['Nota média - Com observabilidade']:.3f}",
    )
    add_metric_triplet(
        slide, x=5.16, y=1.29, w=4.25, title="Conteúdo — NDCG@10",
        initial=f"{content['NDCG - Inicial']:.6f}", attacked=f"{content['NDCG - Sem observabilidade']:.6f}", observed=f"{content['NDCG - Com observabilidade']:.6f}",
    )
    add_round_rect(slide, 0.72, 2.77, 4.02, 1.06, fill=PALE_PURPLE, line=PURPLE)
    add_text(slide, "IMPACTO NO ITEM–ITEM", 0.98, 2.99, 1.62, 0.18, size=8, color=PURPLE, bold=True)
    add_text(slide, f"{item['NDCG - Inicial']:.6f}  →  {item['NDCG - Sem observabilidade']:.6f}  →  {item['NDCG - Com observabilidade']:.6f}", 0.98, 3.32, 3.5, 0.24, size=15, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 5.08, 2.77, 4.2, 1.06, fill=PALE_GREEN, line=GREEN)
    add_text(slide, "RESPOSTA", 5.34, 2.99, 1.0, 0.18, size=8, color=GREEN, bold=True)
    add_text(slide, "A mudança estatística aciona rollback para o snapshot íntegro e restaura os dois modelos.", 5.34, 3.27, 3.68, 0.4, size=10, color=DARK_GREEN, bold=True)
    add_text(slide, "Conteúdo perdeu 19,9% de NDCG; item–item perdeu 100%. Em ambos, a condição observável retornou ao valor inicial.", 0.88, 4.27, 8.24, 0.35, size=11, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: abas Dados e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 65 segundos. Defina drift como mudança de distribuição. A média subiu de 3,822 para 4,182, o conteúdo perdeu 19,9% e o item-item zerou. O rollback recuperou o snapshot e as métricas.")


def _ndcg_table_rows(results, attack: str) -> list[tuple[str, str, str, str]]:
    models = [
        ("Popularidade", "Popularidade regularizada"),
        ("Conteúdo por tags", "Conteúdo por tags"),
        ("Colaborativo item-item", "Colaborativo item-item"),
    ]
    rows = []
    for label, model in models:
        row = measured_row(results, "quality", model, attack)
        rows.append(
            (
                label,
                f"{row['NDCG - Inicial']:.6f}",
                f"{row['NDCG - Sem observabilidade']:.6f}",
                f"{row['NDCG - Com observabilidade']:.6f}",
            )
        )
    return rows


def _coverage_table_rows(results, attack: str) -> list[tuple[str, str, str, str]]:
    models = [
        ("Popularidade", "Popularidade regularizada"),
        ("Conteúdo por tags", "Conteúdo por tags"),
        ("Colaborativo item-item", "Colaborativo item-item"),
    ]
    rows = []
    for label, model in models:
        row = measured_row(results, "quality", model, attack)
        rows.append(
            (
                label,
                f"{row['Cobertura - Inicial']:.4f}",
                f"{row['Cobertura - Sem observabilidade']:.4f}",
                f"{row['Cobertura - Com observabilidade']:.4f}",
            )
        )
    return rows


def slide_push_results_table(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    attack = "Push attack"
    data = measured_row(results, "data", "Popularidade regularizada", attack)
    injected = data["Registros - Sem observabilidade"] - data["Registros - Inicial"]
    add_slide_title(slide, "Resultados obtidos", "Push attack: promoção coordenada do item-alvo")
    add_metric_triplet(
        slide,
        x=0.7,
        y=1.25,
        w=4.2,
        title="Nota média do snapshot",
        initial=f"{data['Nota média - Inicial']:.3f}",
        attacked=f"{data['Nota média - Sem observabilidade']:.3f}",
        observed=f"{data['Nota média - Com observabilidade']:.3f}",
    )
    add_explanation_card(
        slide,
        5.18,
        1.24,
        4.12,
        1.0,
        "O que ocorreu",
        f"{injected:,.0f} interações artificiais promoveram o alvo. A detecção dos novos perfis acionou rollback do lote.".replace(",", "."),
        color=ORANGE,
        fill=PALE_ORANGE,
    )
    add_model_result_table(slide, "Cobertura do catálogo por modelo", _coverage_table_rows(results, attack), y=2.48)
    add_round_rect(slide, 0.72, 4.5, 8.56, 0.5, fill=PALE_BLUE, line=BLUE)
    add_text(
        slide,
        "A cobertura do item-item caiu de 0,0938 para 0,0929 e foi restaurada; os demais modelos não mudaram. A exposição e a posição do alvo ainda são necessárias como métricas diretas.",
        0.95,
        4.62,
        8.1,
        0.28,
        size=9.1,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Fonte: abas Dados, Qualidade e Operação da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 60 segundos. Leia a cobertura por modelo: apenas item-item apresentou queda mensurável e recuperou. Explique que cobertura é uma evidência indireta; posição e exposição do alvo são as métricas específicas que devem ser acrescentadas.")


def slide_nuke_results_table(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    attack = "Nuke attack"
    data = measured_row(results, "data", "Popularidade regularizada", attack)
    injected = data["Registros - Sem observabilidade"] - data["Registros - Inicial"]
    add_slide_title(slide, "Resultados obtidos", "Nuke attack: supressão coordenada do item-alvo")
    add_metric_triplet(
        slide,
        x=0.7,
        y=1.25,
        w=4.2,
        title="Nota média do snapshot",
        initial=f"{data['Nota média - Inicial']:.3f}",
        attacked=f"{data['Nota média - Sem observabilidade']:.3f}",
        observed=f"{data['Nota média - Com observabilidade']:.3f}",
    )
    add_explanation_card(
        slide,
        5.18,
        1.24,
        4.12,
        1.0,
        "Por que a média aumentou?",
        f"O alvo recebeu nota 1, mas os cinco livros de apoio receberam nota 5 em {injected:,.0f} novas interações; por isso a média global subiu.".replace(",", "."),
        color=PURPLE,
        fill=PALE_PURPLE,
    )
    add_model_result_table(slide, "Cobertura do catálogo por modelo", _coverage_table_rows(results, attack), y=2.48)
    add_round_rect(slide, 0.72, 4.5, 8.56, 0.5, fill=PALE_GREEN, line=GREEN)
    add_text(
        slide,
        "A cobertura mudou na popularidade e no item-item e retornou ao baseline após rollback. Para medir a supressão diretamente, ainda são necessárias posição e exposição do alvo.",
        0.95,
        4.62,
        8.1,
        0.28,
        size=9.1,
        color=DARK_GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Fonte: abas Dados, Qualidade e Operação da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 60 segundos. Explique o aparente paradoxo da média: o livro-alvo recebe 1, mas os fillers recebem 5. A tabela mostra as mudanças de cobertura e a recuperação; posição e exposição mediriam diretamente a supressão do alvo.")


def slide_corruption_results_table(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    attack = "Corrupção de dados"
    data = measured_row(results, "data", "Conteúdo por tags", attack)
    add_slide_title(slide, "Resultados obtidos", "Corrupção: impacto dos dados por modelo")
    add_metric_triplet(
        slide,
        x=0.7,
        y=1.25,
        w=4.2,
        title="Score de qualidade dos dados",
        initial=f"{data['Score de qualidade - Inicial']:.3f}",
        attacked=f"{data['Score de qualidade - Sem observabilidade']:.3f}",
        observed=f"{data['Score de qualidade - Com observabilidade']:.3f}",
    )
    add_explanation_card(
        slide,
        5.18,
        1.24,
        4.12,
        1.0,
        "O que ocorreu",
        "O lote chegou a 7,69% de inválidos, 7,69% de órfãos e 37,5% de duplicidade. A resposta colocou o lote em quarentena.",
        color=ORANGE,
        fill=PALE_ORANGE,
    )
    add_model_result_table(slide, "NDCG@10 por modelo", _ndcg_table_rows(results, attack), y=2.48)
    add_round_rect(slide, 0.72, 4.5, 8.56, 0.5, fill=PALE_BLUE, line=BLUE)
    add_text(
        slide,
        "Conteúdo perdeu 9,5% e foi restaurado. Popularidade ficou no piso zero. O aumento isolado do item-item não representa melhoria: o score 0,471 comprova a degradação dos dados.",
        0.95,
        4.62,
        8.1,
        0.28,
        size=9.0,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Fonte: abas Dados e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 65 segundos. Mostre o score de dados e depois leia a tabela: conteúdo caiu e recuperou; popularidade já estava em zero; item-item oscilou para cima, o que não deve ser interpretado como benefício da corrupção.")


def slide_drift_results_table(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    attack = "Drift de distribuição"
    data = measured_row(results, "data", "Conteúdo por tags", attack)
    add_slide_title(slide, "Resultados obtidos", "Drift: mudança de distribuição e impacto por modelo")
    add_metric_triplet(
        slide,
        x=0.7,
        y=1.25,
        w=4.2,
        title="Nota média do snapshot",
        initial=f"{data['Nota média - Inicial']:.3f}",
        attacked=f"{data['Nota média - Sem observabilidade']:.3f}",
        observed=f"{data['Nota média - Com observabilidade']:.3f}",
    )
    add_explanation_card(
        slide,
        5.18,
        1.24,
        4.12,
        1.0,
        "O que ocorreu",
        "Uma fração controlada das avaliações foi deslocada para notas altas. O monitoramento detectou a mudança estatística e aplicou rollback.",
        color=ORANGE,
        fill=PALE_ORANGE,
    )
    add_model_result_table(slide, "NDCG@10 por modelo", _ndcg_table_rows(results, attack), y=2.48)
    add_round_rect(slide, 0.72, 4.5, 8.56, 0.5, fill=PALE_GREEN, line=GREEN)
    add_text(
        slide,
        "Popularidade permaneceu no piso zero; conteúdo perdeu 19,9%; item-item caiu 100%. Após o rollback, conteúdo e item-item retornaram exatamente ao valor inicial.",
        0.95,
        4.62,
        8.1,
        0.28,
        size=9.1,
        color=DARK_GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Fonte: abas Dados e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 65 segundos. A tabela deixa clara a sensibilidade de cada modelo: popularidade não mudou porque já estava no piso; conteúdo sofreu perda parcial; item-item zerou. O rollback restaurou os dois modelos personalizados.")


def slide_outage_results_table(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    attack = "Indisponibilidade"
    operation_rows = [row for row in results["operation"] if row["Ataque"] == attack]
    failures_without = sum(row["Falhas - Sem observabilidade"] for row in operation_rows)
    failures_with = sum(row["Falhas - Com observabilidade"] for row in operation_rows)
    add_slide_title(slide, "Resultados obtidos", "Indisponibilidade: operação × relevância")
    add_metric_triplet(slide, x=0.7, y=1.25, w=4.2, title="Disponibilidade", initial="100", attacked="41", observed="99", unit="%")
    add_explanation_card(
        slide,
        5.18,
        1.24,
        4.12,
        1.0,
        "O que ocorreu",
        f"O fallback por popularidade reduziu as falhas agregadas de {failures_without} para {failures_with}, mas trocou os rankings personalizados por uma lista global.",
        color=ORANGE,
        fill=PALE_ORANGE,
    )
    add_model_result_table(slide, "NDCG@10 por modelo", _ndcg_table_rows(results, attack), y=2.48)
    add_round_rect(slide, 0.72, 4.5, 8.56, 0.5, fill=PALE_BLUE, line=BLUE)
    add_text(
        slide,
        "A observabilidade recuperou +58 p.p. de disponibilidade e reduziu 98,3% das falhas. Porém, o fallback herdou o NDCG zero da popularidade e não preservou a relevância.",
        0.95,
        4.62,
        8.1,
        0.28,
        size=9.0,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Fonte: abas Operação e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 70 segundos. Leia a disponibilidade, depois a tabela de NDCG. Popularidade permanece zero; conteúdo e item-item perdem personalização quando o fallback assume. Este slide explicita o trade-off entre operação e relevância.")


def slide_latency_results(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Latência: resposta parcial, sem perda de relevância")
    attack = "Degradação de latência"
    model_names = ["Popularidade regularizada", "Conteúdo por tags", "Colaborativo item-item"]
    rows = [measured_row(results, "operation", model, attack) for model in model_names]
    initial = mean(row["Latência p95 - Inicial (ms)"] for row in rows)
    attacked = mean(row["Latência p95 - Sem observabilidade (ms)"] for row in rows)
    observed = mean(row["Latência p95 - Com observabilidade (ms)"] for row in rows)
    add_metric_triplet(slide, x=0.7, y=1.3, w=4.2, title="p95 médio dos três modelos", initial=f"{initial:.2f}", attacked=f"{attacked:.2f}", observed=f"{observed:.2f}", unit=" ms")
    add_explanation_card(slide, 5.18, 1.29, 4.12, 0.98, "O que ocorreu", "O atraso elevou o p95 em 147%. Circuit breaker e desvio do caminho degradado reduziram 11,5% da latência atacada.", color=ORANGE, fill=PALE_ORANGE)

    add_text(slide, "p95 por modelo (ms)", 0.72, 2.62, 2.1, 0.22, size=11, color=DARK_GREEN, bold=True)
    headers = ["Modelo", "Inicial", "Sem obs.", "Com obs."]
    widths = [3.35, 1.5, 1.5, 1.5]
    x_positions = [0.72, 4.07, 5.57, 7.07]
    for text_value, x, width in zip(headers, x_positions, widths):
        add_round_rect(slide, x, 2.93, width, 0.38, fill=DARK_GREEN, line=DARK_GREEN, radius=False)
        add_text(slide, text_value, x + 0.05, 3.04, width - 0.1, 0.14, size=7.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    short_names = ["Popularidade", "Conteúdo por tags", "Colaborativo item-item"]
    for index, (short, row) in enumerate(zip(short_names, rows)):
        y = 3.31 + index * 0.4
        values = [short, f"{row['Latência p95 - Inicial (ms)']:.3f}", f"{row['Latência p95 - Sem observabilidade (ms)']:.3f}", f"{row['Latência p95 - Com observabilidade (ms)']:.3f}"]
        for cell_index, (value, x, width) in enumerate(zip(values, x_positions, widths)):
            add_round_rect(slide, x, y, width, 0.38, fill=WHITE if index % 2 == 0 else LIGHT, line=BORDER, radius=False)
            add_text(slide, value, x + 0.05, y + 0.11, width - 0.1, 0.14, size=7.5, color=DARK if cell_index == 0 else TEXT, bold=cell_index == 0, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 0.72, 4.62, 8.56, 0.32, fill=PALE_BLUE, line=PALE_BLUE)
    add_text(slide, "NDCG permaneceu igual: esse ataque degrada tempo de resposta, não a ordem do ranking. A disponibilidade ficou em 100%.", 0.94, 4.7, 8.12, 0.16, size=9.3, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: abas Operação e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 70 segundos. Explique p95 como cauda de latência. A defesa melhorou o cenário, mas não voltou ao baseline; por isso o resultado é parcial. O NDCG inalterado é esperado porque a lista continua igual, apenas mais lenta.")


def slide_outage_results(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Indisponibilidade: disponibilidade × relevância")
    operation_rows = [row for row in results["operation"] if row["Ataque"] == "Indisponibilidade"]
    content = measured_row(results, "quality", "Conteúdo por tags", "Indisponibilidade")
    item = measured_row(results, "quality", "Colaborativo item-item", "Indisponibilidade")
    failures_without = sum(row["Falhas - Sem observabilidade"] for row in operation_rows)
    failures_with = sum(row["Falhas - Com observabilidade"] for row in operation_rows)
    add_metric_triplet(slide, x=0.62, y=1.28, w=4.2, title="Disponibilidade", initial="100", attacked="41", observed="99", unit="%")
    add_metric_triplet(slide, x=5.18, y=1.28, w=4.2, title="Falhas agregadas", initial="0", attacked=str(failures_without), observed=str(failures_with))
    add_round_rect(slide, 0.68, 2.78, 4.08, 1.08, fill=PALE_GREEN, line=GREEN)
    add_text(slide, "GANHO OPERACIONAL", 0.94, 3.0, 1.62, 0.18, size=8, color=GREEN, bold=True)
    add_text(slide, "Fallback por popularidade recuperou +58 pontos percentuais de disponibilidade e reduziu as falhas em 98,3%.", 0.94, 3.31, 3.55, 0.4, size=10.3, color=DARK_GREEN, bold=True)
    add_round_rect(slide, 5.06, 2.78, 4.22, 1.08, fill=PALE_ORANGE, line=ORANGE)
    add_text(slide, "CUSTO NA RELEVÂNCIA", 5.32, 3.0, 1.7, 0.18, size=8, color=ORANGE, bold=True)
    add_text(slide, f"Conteúdo: {content['NDCG - Inicial']:.6f} → {content['NDCG - Com observabilidade']:.6f}\nItem–item: {item['NDCG - Inicial']:.6f} → {item['NDCG - Com observabilidade']:.6f}", 5.32, 3.27, 3.7, 0.48, size=9.6, color=DARK, bold=True)
    add_round_rect(slide, 0.72, 4.23, 8.56, 0.52, fill=PALE_BLUE, line=BLUE)
    add_text(slide, "A observabilidade manteve o serviço respondendo, mas o fallback disponível não preservou a personalização. Disponibilidade e qualidade devem ser acompanhadas juntas.", 0.98, 4.36, 8.04, 0.28, size=9.7, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: abas Operação e Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 70 segundos. Este é o trade-off mais importante: a defesa recupera a operação, mas a lista de popularidade não mantém a relevância dos modelos personalizados. Observabilidade revela essa troca; não a elimina automaticamente.")


def slide_quality_summary(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Resultados obtidos", "Síntese da qualidade de recomendação")
    initial_rows = [
        measured_row(results, "quality", "Popularidade regularizada", "Push attack"),
        measured_row(results, "quality", "Conteúdo por tags", "Push attack"),
        measured_row(results, "quality", "Colaborativo item-item", "Push attack"),
    ]
    models = [
        ("Popularidade", initial_rows[0], ORANGE),
        ("Conteúdo", initial_rows[1], BLUE),
        ("Item–item", initial_rows[2], PURPLE),
    ]
    for index, (label, row, color) in enumerate(models):
        x = 0.67 + index * 3.05
        add_round_rect(slide, x, 1.28, 2.72, 1.22, fill=WHITE, line=color)
        add_text(slide, label, x + 0.16, 1.47, 2.4, 0.22, size=12, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, f"NDCG  {row['NDCG - Inicial']:.6f}", x + 0.16, 1.86, 2.4, 0.22, size=13, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, f"cobertura  {row['Cobertura - Inicial']:.4f}", x + 0.16, 2.18, 2.4, 0.16, size=8, color=MUTED, align=PP_ALIGN.CENTER)

    lower = equal = higher = 0
    for row in results["quality"]:
        delta = row["NDCG - Sem observabilidade"] - row["NDCG - Inicial"]
        if delta < -1e-12:
            lower += 1
        elif delta > 1e-12:
            higher += 1
        else:
            equal += 1
    states = [(str(lower), "quedas", ORANGE), (str(equal), "iguais", BLUE), (str(higher), "aumento isolado", PURPLE)]
    add_text(slide, "NDCG inicial × cenário atacado nas 18 combinações", 0.75, 2.86, 4.2, 0.24, size=11, color=DARK_GREEN, bold=True)
    for index, (value, label, color) in enumerate(states):
        x = 0.76 + index * 1.42
        add_round_rect(slide, x, 3.22, 1.2, 0.76, fill=WHITE, line=color)
        add_text(slide, value, x + 0.08, 3.35, 1.04, 0.23, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.07, 3.68, 1.06, 0.16, size=7, color=MUTED, align=PP_ALIGN.CENTER)
    add_explanation_card(slide, 5.1, 2.82, 4.18, 1.18, "O que os números demonstram", "Corrupção e drift degradaram a qualidade de modelos sensíveis e foram revertidos. Latência e push/nuke não alteraram o NDCG. Outage mostrou conflito entre disponibilidade e relevância.", color=GREEN, fill=PALE_GREEN)
    add_round_rect(slide, 0.72, 4.37, 8.56, 0.45, fill=PALE_BLUE, line=PALE_BLUE)
    add_text(slide, "Um valor maior após corrupção no item–item foi uma oscilação isolada; não supera a evidência direta de dados inválidos e duplicados.", 0.96, 4.49, 8.06, 0.22, size=8.7, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: aba Qualidade da bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 65 segundos. Conteúdo teve maior NDCG inicial e item-item maior cobertura. Mostre a contagem de quedas, empates e uma melhora isolada. Explique que nem todo ataque deve reduzir NDCG e que uma oscilação positiva não torna dados corrompidos desejáveis.")


def slide_validation_measured(prs, results):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Implementação", "Validação funcional e reprodutibilidade")
    meta = results["methodology"]
    stats = [
        ("27/27", "testes backend", GREEN),
        ("4/4", "testes frontend", BLUE),
        (f"{meta['Livros no snapshot']:,}".replace(",", "."), "livros no snapshot", PURPLE),
        (f"{meta['Avaliações no snapshot']:,}".replace(",", "."), "avaliações", ORANGE),
    ]
    for index, (value, label, color) in enumerate(stats):
        x = 0.62 + index * 2.28
        add_round_rect(slide, x, 1.3, 2.04, 0.92, fill=WHITE, line=color)
        add_text(slide, value, x + 0.1, 1.5, 1.84, 0.26, size=18, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.1, 1.89, 1.84, 0.16, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    columns = [
        (0.67, "SISTEMA", ["API e interface integradas", "autenticação e autorização", "carga dos CSVs", "exportação XLSX"]),
        (3.54, "EXPERIMENTO", ["snapshot controlado", "ataques sem alterar o banco", "seed e coorte pareadas", "36 cenários medidos"]),
        (6.41, "ENTREGA", ["Docker local", "manifests Kubernetes/EKS", "deploy e destruição AWS", "relatório reproduzível"]),
    ]
    for x, heading, bullets in columns:
        add_round_rect(slide, x, 2.6, 2.5, 1.75, fill=LIGHT, line=BORDER)
        add_text(slide, heading, x + 0.16, 2.82, 2.18, 0.2, size=9, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        add_bullet_list(slide, bullets, x + 0.2, 3.18, 2.1, 1.0, size=8.5, gap=0.27)
    add_round_rect(slide, 0.78, 4.58, 8.42, 0.35, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "A bateria foi concluída em 102,4 s e gerou quatro abas: Qualidade, Operação, Dados e Metodologia.", 0.98, 4.67, 8.02, 0.18, size=9.3, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fonte: suíte automatizada do repositório e bateria medida em 19/10/2026.")
    add_note(slide, "Tempo sugerido: 55 segundos. Apresente os testes e a escala dos dados. Reforce que o mesmo fluxo pode ser executado localmente e na AWS, e que a saída XLSX registra tanto qualidade quanto operação.")


def slide_interpretation(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Discussão", "Onde a observabilidade ajudou — e onde não bastou")
    rows = [
        ("Corrupção", "Recuperação completa", "quarentena restaurou score e NDCG de conteúdo", GREEN),
        ("Drift", "Recuperação completa", "rollback restaurou distribuição e ranking", GREEN),
        ("Latência", "Melhora parcial", "p95 caiu, mas permaneceu acima do baseline", BLUE),
        ("Outage", "Operação recuperada", "disponibilidade voltou; personalização não", ORANGE),
        ("Push / nuke", "Detecção e rollback", "NDCG não mede diretamente o alvo atacado", PURPLE),
    ]
    for index, (attack, status, explanation, color) in enumerate(rows):
        y = 1.28 + index * 0.68
        add_round_rect(slide, 0.67, y, 8.66, 0.54, fill=WHITE if index % 2 == 0 else LIGHT, line=BORDER)
        add_text(slide, attack, 0.88, y + 0.16, 1.38, 0.18, size=9.5, color=color, bold=True)
        add_text(slide, status, 2.45, y + 0.16, 1.82, 0.18, size=9.5, color=DARK, bold=True)
        add_text(slide, explanation, 4.32, y + 0.16, 4.75, 0.18, size=8.7, color=TEXT)
    add_round_rect(slide, 0.72, 4.82, 8.56, 0.32, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "Síntese: observabilidade reduz exposição e orienta respostas, mas a defesa e a métrica precisam corresponder ao mecanismo da falha.", 0.92, 4.9, 8.16, 0.16, size=9.2, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 55 segundos. Este slide responde diretamente à pergunta de pesquisa. Destaque que a observabilidade teve benefícios diferentes: recuperação completa, parcial, operacional ou apenas diagnóstica.")


def slide_conclusion_measured(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_title(slide, "Conclusões", "Resultados alcançados e próximos refinamentos")
    contributions = [
        ("1", "Protótipo integrado", "Sistema de recomendação, banco, interface, observabilidade e execução local/AWS."),
        ("2", "Evidência experimental", "36 cenários pareados quantificam ranking, dados, latência e disponibilidade."),
        ("3", "Conclusão científica", "A observabilidade melhora a recuperação, mas não substitui defesa específica nem métrica adequada."),
    ]
    for index, (number, title, body) in enumerate(contributions):
        y = 1.32 + index * 1.0
        add_circle(slide, 0.68, y, 0.48, GREEN, number, size=13)
        add_text(slide, title, 1.38, y, 3.72, 0.24, size=13, color=DARK_GREEN, bold=True)
        add_text(slide, body, 1.38, y + 0.34, 3.72, 0.36, size=9.6, color=TEXT)
    add_round_rect(slide, 5.42, 1.3, 3.85, 3.02, fill=LIGHT, line=BORDER)
    add_text(slide, "PRÓXIMOS REFINAMENTOS", 5.72, 1.62, 2.55, 0.2, size=9, color=GREEN, bold=True)
    add_bullet_list(
        slide,
        [
            "medir exposição e posição do item-alvo",
            "repetir seeds e calcular intervalos de confiança",
            "melhorar fallback preservando personalização",
            "completar correlação de logs e traces",
        ],
        5.7, 2.04, 3.12, 1.65, size=9.6, gap=0.38,
    )
    add_round_rect(slide, 0.78, 4.58, 8.42, 0.38, fill=PALE_GREEN, line=PALE_GREEN)
    add_text(slide, "O sistema não apenas recomenda: ele permite demonstrar quando, por que e até que ponto a qualidade se recupera.", 0.98, 4.68, 8.02, 0.18, size=10, color=DARK_GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 60 segundos. Feche com as três entregas: sistema, experimento e conclusão científica. Os próximos refinamentos não anulam os resultados; aumentam a capacidade explicativa e a validade estatística.")


def make_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    add_text(slide, "Obrigado!", 3.15, 4.42, 3.7, 0.42, size=25, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Perguntas?", 3.75, 4.87, 2.5, 0.22, size=12, color=LIME, align=PP_ALIGN.CENTER)
    add_note(slide, "Tempo sugerido: 10 segundos. Agradeça e abra para perguntas.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--assets", required=True, type=Path)
    parser.add_argument("--results", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    results = load_measured_results(args.results)

    prs = Presentation(str(args.template))
    remove_all_slides(prs)
    prs.core_properties.title = "Monitoramento de Qualidade dos Dados para Personalização de Recomendação de Produtos"
    prs.core_properties.author = "Arthur Guarizi Gasparim"
    prs.core_properties.subject = "Apresentação final de iniciação científica"
    prs.core_properties.comments = "Gerado a partir do anteprojeto e do padrão institucional fornecidos."

    make_cover(prs)
    slide_context(prs)
    slide_objective(prs)
    slide_observability(prs)
    slide_metrics_explained(prs)
    slide_recommendation_measured(prs, results)
    slide_data(prs)
    slide_architecture(prs)
    slide_system(prs, args.assets / "sistema_catalogo.png")
    slide_validation_measured(prs, results)
    slide_real_protocol(prs, results)
    slide_attack_map(prs)
    slide_push_results_table(prs, results)
    slide_nuke_results_table(prs, results)
    slide_corruption_results_table(prs, results)
    slide_drift_results_table(prs, results)
    slide_latency_results(prs, results)
    slide_outage_results_table(prs, results)
    slide_quality_summary(prs, results)
    slide_interpretation(prs)
    slide_conclusion_measured(prs)
    make_closing(prs)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.output))
    print(f"Criado: {args.output} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
